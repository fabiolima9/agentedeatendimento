import os
from pathlib import Path
import re
import json

import streamlit as st
from dotenv import load_dotenv
import google.generativeai as genai

# Leitura opcional de PDF e DOCX
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

# -------------------------------------------------------------------
# Configuração e Design Gráfico Executivo (CSS Avançado)
# -------------------------------------------------------------------
PASTA_PROJETO = Path(__file__).resolve().parent
ARQUIVO_ENV = PASTA_PROJETO / ".env"
load_dotenv(dotenv_path=ARQUIVO_ENV, override=False)

st.set_page_config(
    page_title="OrientAção IA - Plataforma de Carreira",
    page_icon="🧭",
    layout="wide",
)

# Injeção de CSS personalizado para UI/UX profissional (SaaS / Enterprise)
st.markdown("""
    <style>
    .main {
        background-color: #f8fafc;
        font-family: 'Segoe UI', -apple-system, BlinkMacSystemFont, Roboto, Helvetica, Arial, sans-serif;
    }
    .stSidebar {
        background-color: #f1f5f9;
        border-right: 1px solid #e2e8f0;
    }
    h1, h2, h3 {
        color: #0f172a;
        font-weight: 700;
        letter-spacing: -0.5px;
    }
    .stButton>button {
        border-radius: 8px;
        font-weight: 600;
        transition: all 0.2s ease-in-out;
        border: 1px solid #cbd5e1;
        background-color: #ffffff;
        color: #0f172a;
    }
    .stButton>button:hover {
        border-color: #0284c7;
        color: #0284c7;
        box-shadow: 0 4px 12px rgba(2, 132, 199, 0.15);
    }
    .stFormSubmitButton>button {
        background-color: #0f172a;
        color: #ffffff;
        border: none;
    }
    .stFormSubmitButton>button:hover {
        background-color: #1e293b;
        color: #ffffff;
    }
    div.stAlert {
        border-radius: 10px;
        border: 1px solid #e2e8f0;
    }
    </style>
""", unsafe_allow_html=True)

MODOS = [
    "🧭 Orientação profissional",
    "📊 Análise de competências",
    "📈 Plano de carreira",
    "📄 Currículo profissional (Análise & Otimização)",
    "🔗 Análise de vaga & Portais de Emprego",
    "💬 Simulação de entrevista",
    "🎙️ Simulação de Entrevista por Voz",
    "😈 Advogado do Diabo (Desafiador Incisivo)",
    "📝 Teste Prático de Conhecimento",
    "⚖️ Módulo de Pressão Comportamental",
    "👑 Módulo de Liderança Executiva",
    "🧠 Comportamento e Postura Profissional",
    "📊 Gerar Apresentação de Slides (Projeto)",
]

SYSTEM_PROMPT = (
    "Você é o OrientAção IA, um assistente executivo de orientação profissional, desenvolvimento "
    "de carreira, empregabilidade e mentoria comportamental no Brasil. "
    "Responda em português brasileiro, com tom altamente profissional, prático e respeitoso.\n\n"
    "REGRAS DE CONFIABILIDADE E MENTORIA:\n"
    "- Use apenas os dados que a pessoa forneceu como fatos sobre ela.\n"
    "- Diferencie claramente 'informado pela pessoa', 'possibilidade a confirmar'.\n"
    "- Se os dados forem insuficientes, faça até 3 perguntas objetivas.\n"
    "- MENTORIA COMPORTAMENTAL: Ao avaliar conflitos ou postura, utilize princípios de Inteligência Emocional e CNV.\n"
    "- CORREÇÃO DE POSTURA: Identifique atitudes que sabotam o ambiente de trabalho e explique o impacto negativo.\n"
    "- Seja honesto e direto ao apontar erros, sempre de forma construtiva.\n"
    "- Use Markdown limpo e legível. Evite recomendações genéricas."
)

def limpar_resposta(texto: str) -> str:
    texto = re.sub(r"(?i)<br\s*/?>", "\n", texto)
    texto = re.sub(r"<[^>]+>", "", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    return texto.strip()

def extrair_curriculo(arquivo) -> tuple[str, str | None]:
    if arquivo is None:
        return "", None
    nome = arquivo.name.lower()
    try:
        if nome.endswith(".txt"):
            conteudo = arquivo.getvalue()
            for encoding in ("utf-8", "latin-1"):
                try:
                    return conteudo.decode(encoding).strip(), None
                except UnicodeDecodeError:
                    continue
            return "", "Não consegui ler o arquivo TXT."
        
        if nome.endswith(".pdf"):
            if PdfReader is None:
                return "", "Para ler PDF, instale: pip install pypdf"
            reader = PdfReader(arquivo)
            partes = [pagina.extract_text() for pagina in reader.pages if pagina.extract_text()]
            texto = "\n".join(partes).strip()
            return (texto, None) if texto else ("", "Não encontrei texto selecionável nesse PDF.")
        
        if nome.endswith(".docx"):
            if Document is None:
                return "", "Para ler DOCX, instale: pip install python-docx"
            documento = Document(arquivo)
            partes = [p.text for p in documento.paragraphs if p.text.strip()]
            for tabela in documento.tables:
                for linha in tabela.rows:
                    celulas = [c.text.strip() for c in linha.cells]
                    partes.append(" | ".join(celulas))
            texto = "\n".join(partes).strip()
            return (texto, None) if texto else ("", "O arquivo DOCX parece estar vazio.")
            
        return "", "Formato não reconhecido. Use PDF, DOCX ou TXT."
    except Exception as exc:
        return "", f"Erro ao ler arquivo: {type(exc).__name__}"

def parece_analise_de_curriculo(mensagem: str) -> bool:
    texto = mensagem.lower()
    termos = ("analise meu curriculo", "analisa meu curriculo", "analisar meu curriculo", "revisar meu curriculo", "revise meu curriculo", "corrigir meu curriculo", "avalie meu curriculo", "avaliar meu curriculo", "melhore meu curriculo")
    normalizado = texto.replace("á", "a").replace("à", "a").replace("ã", "a").replace("â", "a").replace("é", "e").replace("ê", "e").replace("í", "i").replace("ó", "o").replace("õ", "o")
    return any(termo in normalizado for termo in termos)

# -------------------------------------------------------------------
# Interface Principal
# -------------------------------------------------------------------
col_logo, col_title = st.columns([1, 8])
with col_logo:
    st.markdown("# 🧭")
with col_title:
    st.title("OrientAção IA")
    st.markdown("##### *Plataforma Executiva de Carreira, Inteligência Comportamental e Empregabilidade*")

st.markdown("---")

with st.sidebar:
    st.markdown("### 👤 Perfil Profissional")
    st.markdown("<small style='color: #64748b;'>Configure seus dados para personalizar as análises e mentorias da IA.</small>", unsafe_allow_html=True)
    
    with st.form("form_perfil"):
        experiencia = st.text_area("Experiência Principal", placeholder="Ex: manutenção, atendimento, suporte...", height=90)
        formacao = st.text_area("Formação & Certificações", placeholder="Ex: técnico, superior, cursos...", height=75)
        competencias = st.text_area("Competências (Hard & Soft Skills)", placeholder="Ex: elétrica, hidráulica, Excel...", height=75)
        objetivo = st.text_input("Objetivo Profissional", placeholder="Ex: promoção, recolocação")
        salario = st.text_input("Salário Desejado (opcional)")
        disponibilidade = st.text_input("Disponibilidade (opcional)")
        salvar_perfil = st.form_submit_button("💾 Salvar Perfil no Sistema", use_container_width=True)

    if salvar_perfil:
        st.session_state.perfil = {
            "Experiência": experiencia.strip(),
            "Formação e cursos": formacao.strip(),
            "Competências informadas": competencias.strip(),
            "Objetivo profissional": objetivo.strip(),
            "Salário desejado": salario.strip(),
            "Disponibilidade": disponibilidade.strip(),
        }
        st.success("Perfil atualizado com sucesso!")

    st.divider()
    modo = st.selectbox("🎯 Selecione a Ferramenta / Módulo", MODOS, index=0)
    
    st.markdown("<br>", unsafe_allow_html=True)
    if st.button("🗑️ Limpar Sessão / Conversa", use_container_width=True):
        st.session_state.mensagens_gemini = []
        st.session_state.mensagens_exibicao = []
        for k in ["quiz_perguntas", "pressao_iniciada", "lideranca_iniciada", "curriculo_otimizado", "diabo_iniciado"]:
            if k in st.session_state:
                del st.session_state[k]
        st.rerun()

if "perfil" not in st.session_state:
    st.session_state.perfil = {}
if "mensagens_gemini" not in st.session_state:
    st.session_state.mensagens_gemini = []
if "mensagens_exibicao" not in st.session_state:
    st.session_state.mensagens_exibicao = []
if "curriculo_texto" not in st.session_state:
    st.session_state.curriculo_texto = ""

perfil = st.session_state.perfil
perfil_formatado = "\n".join(f"- {chave}: {valor if valor else 'Não informado'}" for chave, valor in perfil.items())

# -------------------------------------------------------------------
# MODO: GERADOR DE SLIDES DO PROJETO
# -------------------------------------------------------------------
if modo == "📊 Gerar Apresentação de Slides (Projeto)":
    st.subheader("📊 Gerador de Apresentação Executiva do Projeto")
    st.markdown("Esta ferramenta compila a arquitetura, o propósito e os diferenciais do **OrientAção IA** em uma apresentação de slides profissional em HTML para você entregar e apresentar para a sua professora.")
    
    if st.button("🚀 Gerar Apresentação do Projeto Agora", use_container_width=True):
        api_key = os.getenv("GEMINI_API_KEY", "").strip()
        if not api_key:
            st.error("⚠️ Chave `GEMINI_API_KEY` não encontrada no arquivo `.env`.")
        else:
            with st.spinner("Compilando dados do projeto e gerando os slides..."):
                try:
                    genai.configure(api_key=api_key)
                    model = genai.GenerativeModel('gemini-3.1-flash-lite')
                    
                    prompt_slides = (
                        "Gere estritamente um documento HTML completo e válido, iniciando exatamente com <!DOCTYPE html> e terminando com </html>.\n"
                        "O conteúdo deve ser um pitch deck executivo sobre o software 'OrientAção IA' com os seguintes slides:\n"
                        "1. Capa: OrientAção IA - Plataforma de Carreira e Postura.\n"
                        "2. O Problema: Dificuldade de recolocação e falhas de soft skills.\n"
                        "3. A Solução: Mentoria por IA (Google Gemini) focada em inteligência emocional e planos de carreira.\n"
                        "4. Arquitetura: Python, Streamlit, API do Google Gemini, suporte a PDF/DOCX.\n"
                        "5. Diferenciais: Confiabilidade de dados e mentoria comportamental.\n"
                        "Use um design moderno com estilos CSS embutidos (fundo claro, cartões brancos e cores corporativas)."
                    )
                    response = model.generate_content(prompt_slides)
                    html_resultado = limpar_resposta(response.text)
                    html_resultado = re.sub(r"^```html\s*", "", html_resultado, flags=re.IGNORECASE)
                    html_resultado = re.sub(r"^```\s*", "", html_resultado)
                    html_resultado = re.sub(r"\s*```$", "", html_resultado)
                    
                    st.success("✨ Apresentação do projeto gerada com sucesso!")
                    st.download_button(
                        label="📥 Baixar Apresentação em HTML (Abrir no Navegador para Apresentar)",
                        data=html_resultado,
                        file_name="apresentacao_orientacao_ia.html",
                        mime="text/html",
                        use_container_width=True
                    )
                    with st.expander("Pré-visualizar código HTML"):
                        st.code(html_resultado, language="html")
                except Exception as exc:
                    st.error(f"Erro ao gerar slides: {exc}")

# -------------------------------------------------------------------
# MODO: CURRÍCULO PROFISSIONAL
# -------------------------------------------------------------------
elif modo == "📄 Currículo profissional (Análise & Otimização)":
    st.subheader("📄 Módulo de Currículo Profissional: Análise e Otimização Avançada")
    st.markdown("Envie seu currículo ou cole o texto abaixo. A IA fará uma análise crítica de impacto, reescreverá o conteúdo deixando-o pronto e estruturado para você baixar e utilizar nos processos seletivos.")

    arquivo_curriculo = st.file_uploader("Anexar Currículo (PDF, DOCX ou TXT)", type=["pdf", "docx", "txt"])
    if arquivo_curriculo is not None:
        texto_extraido, erro = extrair_curriculo(arquivo_curriculo)
        if erro:
            st.warning(erro)
        elif texto_extraido:
            st.session_state.curriculo_texto = texto_extraido
            st.success("Currículo carregado com sucesso na memória!")

    texto_colado = st.text_area("Ou cole o texto do seu currículo aqui:", value=st.session_state.curriculo_texto, height=200)
    if texto_colado.strip():
        st.session_state.curriculo_texto = texto_colado.strip()

    if st.button("✨ Analisar e Otimizar Currículo Profissional", use_container_width=True):
        if not st.session_state.curriculo_texto.strip():
            st.warning("⚠️ Por favor, envie ou cole o texto do seu currículo primeiro.")
        else:
            api_key = os.getenv("GEMINI_API_KEY", "").strip()
            if not api_key:
                st.error("⚠️ Chave `GEMINI_API_KEY` não encontrada.")
            else:
                with st.spinner("Analisando pontos fracos e reestruturando seu currículo para alto impacto..."):
                    try:
                        genai.configure(api_key=api_key)
                        model = genai.GenerativeModel('gemini-3.1-flash-lite')
                        
                        prompt_cv = (
                            "Analise criticamente o currículo abaixo e reescreva-o completamente de forma profissional, moderna e otimizada (formato pronto para triagem ATS).\n"
                            "Divida sua resposta em duas partes claras:\n"
                            "1. **Diagnóstico e Pontos de Melhoria:** O que estava fraco e o que foi corrigido.\n"
                            "2. **Currículo Otimizado e Pronto:** O texto estruturado do currículo reformulado (com Resumo Profissional forte, Experiências orientadas a resultados e Formação organizada).\n\n"
                            f"CURRÍCULO ORIGINAL:\n{st.session_state.curriculo_texto}"
                        )
                        response = model.generate_content(prompt_cv)
                        st.session_state.curriculo_otimizado = limpar_resposta(response.text)
                        st.success("Currículo otimizado com sucesso!")
                    except Exception as e:
                        st.error(f"Erro ao processar currículo: {e}")

    if "curriculo_otimizado" in st.session_state:
        st.divider()
        st.markdown("### 📋 Resultado da Otimização")
        st.markdown(st.session_state.curriculo_otimizado)
        
        st.download_button(
            label="📥 Baixar Currículo Otimizado (Pronto para Salvar/Imprimir)",
            data=st.session_state.curriculo_otimizado,
            file_name="curriculo_otimizado_orientacao_ia.md",
            mime="text/markdown",
            use_container_width=True
        )

# -------------------------------------------------------------------
# MODO: ANÁLISE DE VAGA & PORTAIS DE EMPREGO
# -------------------------------------------------------------------
elif modo == "🔗 Análise de vaga & Portais de Emprego":
    st.subheader("🔗 Análise de Vaga e Conexão com Portais de Emprego no Brasil")
    st.markdown("Cole a descrição da vaga que você deseja concorrer. A IA fará um cruzamento com o seu perfil cadastrado para apontar aderência, gaps e estratégias de entrevista.")

    st.markdown(
        "### 🌐 Principais Sites de Emprego no Brasil para Acompanhar:\n"
        "- **[LinkedIn](https://www.linkedin.com)** - A principal rede profissional e de recrutamento do mundo.\n"
        "- **[Gupy](https://www.gupy.io)** - Plataforma utilizada pela maioria das grandes empresas brasileiras para triagem automatizada.\n"
        "- **[Catho](https://www.catho.com.br)** - Um dos sites mais tradicionais de vagas e candidaturas no país.\n"
        "- **[Indeed Brasil](https://br.indeed.com)** - Buscador abrangente com milhares de oportunidades locais e remotas.\n"
        "- **[Vagas.com](https://www.vagas.com.br)** - Portal corporativo robusto para posições efetivas e estágios."
    )
    st.divider()

    descricao_vaga = st.text_area("Cole a descrição da vaga aqui:", placeholder="Ex: Requisitos, responsabilidades, conhecimentos exigidos...", height=180)

    if st.button("Analisar Aderência à Vaga", use_container_width=True):
        if not descricao_vaga.strip():
            st.warning("⚠️ Cole a descrição da vaga para realizar a análise.")
        else:
            api_key = os.getenv("GEMINI_API_KEY", "").strip()
            if not api_key:
                st.error("⚠️ Chave `GEMINI_API_KEY` não encontrada.")
            else:
                with st.spinner("Analisando compatibilidade da vaga com seu perfil profissional..."):
                    try:
                        genai.configure(api_key=api_key)
                        model = genai.GenerativeModel('gemini-3.1-flash-lite')
                        
                        prompt_vaga = (
                            "Faça uma análise rigorosa de aderência entre o perfil do usuário e a vaga descrita abaixo.\n\n"
                            f"PERFIL DO USUÁRIO:\n{perfil_formatado}\n\n"
                            f"DESCRIÇÃO DA VAGA:\n{descricao_vaga}\n\n"
                            "Aponte:\n"
                            "1. Taxa estimada de compatibilidade (%).\n"
                            "2. Pontos fortes do usuário que atendem à vaga.\n"
                            "3. Gaps (o que falta) e como compensar no currículo/entrevista.\n"
                            "4. Dicas estratégicas para se destacar nos portais (como Gupy e LinkedIn) para esta oportunidade."
                        )
                        resp = model.generate_content(prompt_vaga)
                        st.markdown(limpar_resposta(resp.text))
                    except Exception as e:
                        st.error(f"Erro na análise: {e}")

# -------------------------------------------------------------------
# MODO: ADVOGADO DO DIABO (DESAFIADOR INCISIVO)
# -------------------------------------------------------------------
elif modo == "😈 Advogado do Diabo (Desafiador Incisivo)":
    st.subheader("😈 Módulo: Advogado do Diabo (Desafiador Incisivo)")
    st.markdown("Neste modo, a IA assume uma postura **implacável, crítica e questionadora**. O objetivo é achar falhas na sua argumentação, testar suas convicções profissionais e expor pontos cegos pessoais para treinar sua resiliência e preparo sob fogo cruzado.")
    
    st.info("🎙️ **Dica de Áudio:** Utilize o microfone do seu teclado (Windows + H no Windows, ou ditado por voz no seu dispositivo) na caixa de chat abaixo para responder falando por áudio!")

    if not st.session_state.get("diabo_iniciado"):
        if st.button("🔥 Iniciar Confronto com o Advogado do Diabo", use_container_width=True):
            api_key = os.getenv("GEMINI_API_KEY", "").strip()
            if not api_key:
                st.error("⚠️ Chave `GEMINI_API_KEY` não encontrada.")
            else:
                with st.spinner("Preparando artilharia crítica e analisando seu perfil..."):
                    genai.configure(api_key=api_key)
                    model = genai.GenerativeModel('gemini-3.1-flash-lite')
                    
                    prompt_diabo_init = (
                        "Inicie uma sessão implacável como 'Advogado do Diabo'. "
                        "Analise o perfil do usuário abaixo e faça o primeiro questionamento duro, incisivo e direto, "
                        "desafiando uma suposição, objetivo ou fraqueza aparente nas competências ou histórico dele:\n\n"
                        f"PERFIL DO USUÁRIO:\n{perfil_formatado}"
                    )
                    resp = model.generate_content(prompt_diabo_init)
                    st.session_state.diabo_chat = [{"role": "assistant", "content": limpar_resposta(resp.text)}]
                    st.session_state.diabo_iniciado = True
                    st.rerun()

    if st.session_state.get("diabo_iniciado"):
        for m in st.session_state.diabo_chat:
            with st.chat_message(m["role"]):
                st.markdown(m["content"])

        resposta_diabo = st.chat_input("Responda ao desafio (digite ou use o ditado por voz)...")
        if resposta_diabo:
            st.session_state.diabo_chat.append({"role": "user", "content": resposta_diabo})
            with st.chat_message("user"):
                st.markdown(resposta_diabo)

            with st.chat_message("assistant"):
                with st.spinner("Analisando falhas e contra-atacando..."):
                    try:
                        genai.configure(api_key=os.getenv("GEMINI_API_KEY").strip())
                        model = genai.GenerativeModel('gemini-3.1-flash-lite')
                        hist_diabo = "\n".join([f"{msg['role']}: {msg['content']}" for msg in st.session_state.diabo_chat])
                        
                        prompt_ataque = (
                            "Você é o Advogado do Diabo. O usuário acabou de responder ao seu questionamento anterior. "
                            "Aponte falhas lógicas, contradições, ingenuidades ou pontos cegos pessoais na resposta dele de forma firme e crítica. "
                            "Faça uma nova provocação ainda mais incisiva para testar os limites dele:\n\n"
                            f"{hist_diabo}"
                        )
                        r = model.generate_content(prompt_ataque)
                        texto_resp = limpar_resposta(r.text)
                        st.markdown(texto_resp)
                        st.session_state.diabo_chat.append({"role": "assistant", "content": texto_resp})
                    except Exception as e:
                        st.error(f"Erro: {e}")

# -------------------------------------------------------------------
# MODO: TESTE PRÁTICO DE CONHECIMENTO
# -------------------------------------------------------------------
elif modo == "📝 Teste Prático de Conhecimento":
    st.subheader("📝 Teste Prático de Competências e Conhecimento")
    st.markdown("Informe abaixo a competência técnica que você deseja testar para receber questões práticas personalizadas geradas pela IA.")

    competencia_alvo = st.text_input("Competência / Assunto:", placeholder="Ex: manutenção elétrica, excel avançado, gestão de frotas")

    if st.button("Gerar Teste Prático Personalizado", use_container_width=True):
        if not competencia_alvo.strip():
            st.warning("⚠️ Por favor, informe a competência.")
        else:
            api_key = os.getenv("GEMINI_API_KEY", "").strip()
            if not api_key:
                st.error("⚠️ Chave `GEMINI_API_KEY` não encontrada.")
            else:
                with st.spinner(f"Elaborando questões práticas sobre '{competencia_alvo}'..."):
                    try:
                        genai.configure(api_key=api_key)
                        model = genai.GenerativeModel('gemini-3.1-flash-lite')
                        
                        prompt_quiz = (
                            f"Crie um teste técnico prático com 3 perguntas de múltipla escolha (A, B, C, D) focadas na competência: '{competencia_alvo}'.\n"
                            f"Perfil do usuário: {perfil_formatado}\n\n"
                            "Retorne EXTREMAMENTE ESTRITAMENTE em formato JSON puro (uma lista de objetos JSON), sem markdown ou textos adicionais:\n"
                            "[\n"
                            "  {\n"
                            "    \"pergunta\": \"...\",\n"
                            "    \"opcoes\": [\"A) ...\", \"B) ...\", \"C) ...\", \"D) ...\"],\n"
                            "    \"resposta_correta\": \"A\",\n"
                            "    \"explicacao\": \"...\"\n"
                            "  }\n"
                            "]"
                        )
                        response = model.generate_content(prompt_quiz)
                        texto_resp = limpar_resposta(response.text)
                        match_json = re.search(r'\[\s*\{.*\}\s*\]', texto_resp, re.DOTALL)
                        if match_json:
                            st.session_state.quiz_perguntas = json.loads(match_json.group(0))
                            st.success(f"Teste de '{competencia_alvo}' gerado com sucesso!")
                        else:
                            st.error("Erro ao estruturar o quiz. Tente novamente.")
                    except Exception as exc:
                        st.error(f"Erro: {exc}")

    if "quiz_perguntas" in st.session_state:
        with st.form("form_quiz"):
            respostas_usuario = {}
            for i, q in enumerate(st.session_state.quiz_perguntas):
                st.markdown(f"**Questão {i+1}:** {q['pergunta']}")
                respostas_usuario[i] = st.radio(f"Opções ({i+1}):", q['opcoes'], key=f"q_{i}")
                st.divider()
            
            if st.form_submit_button("Corrigir Prova", use_container_width=True):
                acertos = sum(1 for i, q in enumerate(st.session_state.quiz_perguntas) if respostas_usuario[i][0] == q['resposta_correta'])
                for i, q in enumerate(st.session_state.quiz_perguntas):
                    if respostas_usuario[i][0] == q['resposta_correta']:
                        st.success(f"Questão {i+1}: Correta! 🎯 {q['explicacao']}")
                    else:
                        st.error(f"Questão {i+1}: Incorreta. Resposta certa: ({q['resposta_correta']}). {q['explicacao']}")
                st.metric("Pontuação Final", f"{acertos} / {len(st.session_state.quiz_perguntas)}")

# -------------------------------------------------------------------
# MODO: MÓDULO DE PRESSÃO COMPORTAMENTAL
# -------------------------------------------------------------------
elif modo == "⚖️ Módulo de Pressão Comportamental":
    st.subheader("⚖️ Simulador de Pressão e Resiliência Corporativa")
    st.markdown("Simule situações de crise ou alta pressão no ambiente de trabalho para treinar inteligência emocional e comunicação assertiva.")

    cenario_usuario = st.text_input("Informe o cenário ou tipo de pressão que deseja simular:", placeholder="Ex: Cobrança agressiva de prazo de entrega por um chefe estressado")

    if st.button("Iniciar Dinâmica de Pressão", use_container_width=True):
        if not cenario_usuario.strip():
            st.warning("⚠️ Descreva um cenário para iniciar a simulação.")
        else:
            api_key = os.getenv("GEMINI_API_KEY", "").strip()
            if not api_key:
                st.error("⚠️ Chave `GEMINI_API_KEY` não encontrada.")
            else:
                with st.spinner("Configurando cenário de pressão..."):
                    genai.configure(api_key=api_key)
                    model = genai.GenerativeModel('gemini-3.1-flash-lite')
                    
                    prompt_inicial = (
                        f"Inicie uma dinâmica interativa de teste de pressão comportamental com base no cenário: '{cenario_usuario}'.\n"
                        f"Perfil profissional: {perfil_formatado}\n"
                        "Aja como um avaliador comportamental rigoroso. Apresente o cenário de crise detalhado e faça uma pergunta direta sobre como o usuário reagiria sob pressão."
                    )
                    resp = model.generate_content(prompt_inicial)
                    st.session_state.pressao_chat = [{"role": "assistant", "content": limpar_resposta(resp.text)}]
                    st.session_state.pressao_iniciada = True
                    st.rerun()

    if st.session_state.get("pressao_iniciada"):
        for m in st.session_state.pressao_chat:
            with st.chat_message(m["role"]):
                st.markdown(m["content"])

        resposta_pressao = st.chat_input("Digite sua reação ou resposta para esta situação de pressão...")
        if resposta_pressao:
            st.session_state.pressao_chat.append({"role": "user", "content": resposta_pressao})
            with st.chat_message("user"):
                st.markdown(resposta_pressao)

            with st.chat_message("assistant"):
                with st.spinner("Analisando sua inteligência emocional e postura..."):
                    try:
                        genai.configure(api_key=os.getenv("GEMINI_API_KEY").strip())
                        model = genai.GenerativeModel('gemini-3.1-flash-lite')
                        hist_pressao = "\n".join([f"{msg['role']}: {msg['content']}" for msg in st.session_state.pressao_chat])
                        
                        prompt_analise = (
                            "Com base no histórico da simulação de pressão abaixo, analise a resposta emocional e profissional do usuário "
                            "sob a ótica da Inteligência Emocional e CNV. Dê o próximo passo do cenário ou nova pergunta:\n"
                            f"{hist_pressao}"
                        )
                        r = model.generate_content(prompt_analise)
                        texto_resp = limpar_resposta(r.text)
                        st.markdown(texto_resp)
                        st.session_state.pressao_chat.append({"role": "assistant", "content": texto_resp})
                    except Exception as e:
                        st.error(f"Erro: {e}")

# -------------------------------------------------------------------
# MODO: MÓDULO DE LIDERANÇA EXECUTIVA
# -------------------------------------------------------------------
elif modo == "👑 Módulo de Liderança Executiva":
    st.subheader("👑 Simulador de Liderança Situacional e Gestão")
    st.markdown("Treine sua capacidade de liderar equipes, delegar tarefas e resolver conflitos com base em desafios reais informados por você.")

    desafio_lideranca = st.text_input("Qual desafio de liderança você deseja enfrentar?", placeholder="Ex: Funcionário desmotivado que não entrega os prazos no setor")

    if st.button("Iniciar Dinâmica de Liderança", use_container_width=True):
        if not desafio_lideranca.strip():
            st.warning("⚠️ Informe um desafio de liderança.")
        else:
            api_key = os.getenv("GEMINI_API_KEY", "").strip()
            if not api_key:
                st.error("⚠️ Chave `GEMINI_API_KEY` não encontrada.")
            else:
                with st.spinner("Preparando estudo de caso de liderança..."):
                    genai.configure(api_key=api_key)
                    model = genai.GenerativeModel('gemini-3.1-flash-lite')
                    
                    prompt_lider_init = (
                        f"Inicie uma simulação prática de liderança executiva com base no desafio: '{desafio_lideranca}'.\n"
                        f"Perfil do líder (usuário): {perfil_formatado}\n"
                        "Aja como um consultor executivo sênior. Apresente o estudo de caso e pergunte qual é a primeira atitude de gestão que o usuário tomaria."
                    )
                    resp = model.generate_content(prompt_lider_init)
                    st.session_state.lideranca_chat = [{"role": "assistant", "content": limpar_resposta(resp.text)}]
                    st.session_state.lideranca_iniciada = True
                    st.rerun()

    if st.session_state.get("lideranca_iniciada"):
        for m in st.session_state.lideranca_chat:
            with st.chat_message(m["role"]):
                st.markdown(m["content"])

        resposta_lider = st.chat_input("Digite sua decisão de gestão para este cenário...")
        if resposta_lider:
            st.session_state.lideranca_chat.append({"role": "user", "content": resposta_lider})
            with st.chat_message("user"):
                st.markdown(resposta_lider)

            with st.chat_message("assistant"):
                with st.spinner("Avaliando sua estratégia de liderança..."):
                    try:
                        genai.configure(api_key=os.getenv("GEMINI_API_KEY").strip())
                        model = genai.GenerativeModel('gemini-3.1-flash-lite')
                        hist_lider = "\n".join([f"{msg['role']}: {msg['content']}" for msg in st.session_state.lideranca_chat])
                        
                        prompt_aval_lider = (
                            "Com base no histórico da simulação de liderança abaixo, avalie a decisão estratégica de gestão do usuário "
                            "e avance o caso com um desdobramento:\n"
                            f"{hist_lider}"
                        )
                        r = model.generate_content(prompt_aval_lider)
                        texto_resp = limpar_resposta(r.text)
                        st.markdown(texto_resp)
                        st.session_state.lideranca_chat.append({"role": "assistant", "content": texto_resp})
                    except Exception as e:
                        st.error(f"Erro: {e}")

# -------------------------------------------------------------------
# OUTROS MÓDULOS DE CHAT
# -------------------------------------------------------------------
else:
    for msg in st.session_state.mensagens_exibicao:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])

    if modo in ["🎙️ Simulação de Entrevista por Voz", "😈 Advogado do Diabo (Desafiador Incisivo)"]:
        st.info("💡 **Dica de Voz:** Utilize o microfone embutido do seu teclado (Windows + H no Windows, ou ditado por voz no celular/PC) na caixa de texto abaixo para conversar falando por áudio!")

    mensagem_usuario = st.chat_input("Escreva sua resposta, pergunta ou simulação...")

    if mensagem_usuario:
        st.session_state.mensagens_exibicao.append({"role": "user", "content": mensagem_usuario})
        with st.chat_message("user"):
            st.markdown(mensagem_usuario)

        api_key = os.getenv("GEMINI_API_KEY", "").strip()
        if not api_key:
            resposta = "⚠️ Chave `GEMINI_API_KEY` não encontrada no arquivo `.env` ou no Render."
            st.session_state.mensagens_exibicao.append({"role": "assistant", "content": resposta})
            with st.chat_message("assistant"):
                st.markdown(resposta)
        else:
            contexto = f"MODO ATUAL: {modo}\n\nPERFIL DO USUÁRIO:\n{perfil_formatado}\n"
            
            modo_chave = modo.split(" ", 1)[1] if " " in modo else modo
            
            instrucoes_modo = {
                "Orientação profissional": "Ajude a pessoa a definir próximos passos práticos de carreira com base nos dados informados.",
                "Análise de competências": "Separe claramente as competências informadas em Hard Skills (técnicas) e Soft Skills (comportamentais), identificando o que precisa ser validado.",
                "Plano de carreira": "Crie um plano prático, estruturado em etapas claras com prazos de curto, médio e longo prazo.",
                "Simulação de entrevista": "Faça uma pergunta por vez de entrevista e dê feedback imediato sobre a resposta.",
                "Simulação de Entrevista por Voz": "Atue como um recrutador exigente em uma entrevista oral. Faça perguntas focadas em comportamento e competência, uma por vez.",
                "Comportamento e Postura Profissional": "Avalie sob a ótica da inteligência emocional. Aponte atitudes ruins que sabotam o ambiente e ensine a agir de forma madura."
            }.get(modo_chave, "Ajude o usuário em sua trajetória profissional com base no perfil fornecido.")

            prompt_completo = f"{SYSTEM_PROMPT}\n\n{contexto}\n\nINSTRUÇÕES DO MÓDULO:\n{instrucoes_modo}\n\nInteração do usuário: {mensagem_usuario}"

            with st.chat_message("assistant"):
                with st.spinner("Processando resposta profissional..."):
                    try:
                        genai.configure(api_key=api_key)
                        model = genai.GenerativeModel('gemini-3.1-flash-lite') 
                        
                        if not st.session_state.mensagens_gemini:
                            chat = model.start_chat(history=[])
                        else:
                            chat = model.start_chat(history=st.session_state.mensagens_gemini)

                        response = chat.send_message(prompt_completo)
                        resposta = limpar_resposta(response.text)
                        st.session_state.mensagens_gemini = chat.history
                    except Exception as exc:
                        resposta = f"❌ **Erro técnico:** `{type(exc).__name__}: {str(exc)}`"
                
                st.markdown(resposta)
            st.session_state.mensagens_exibicao.append({"role": "assistant", "content": resposta})
